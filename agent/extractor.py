"""
extractor.py
------------
Reads FNOL documents (.txt/.pdf/.docx) and images (.jpg/.jpeg/.png/.webp)
and extracts known fields with a label-aware, deterministic parser.
"""
import re
from pathlib import Path

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from PIL import Image
except ImportError:
    Image = None

try:
    from rapidocr_onnxruntime import RapidOCR
except ImportError:
    RapidOCR = None


FIELD_SCHEMA = {
    "policyInformation": {
        "policyNumber": ["Policy Number", "Policy No", "Policy #", "Policy ID"],
        "policyholderName": ["Policyholder Name", "Policy Holder Name", "Name of Insured", "Insured Name"],
        "effectiveDates": ["Effective Dates", "Policy Effective Dates", "Policy Period", "Policy Dates"],
    },
    "incidentInformation": {
        "date": ["Date of Loss", "Incident Date", "Loss Date", "Date of Incident"],
        "time": ["Time of Loss", "Incident Time", "Loss Time", "Time of Incident"],
        "location": ["Location of Loss", "Incident Location", "Loss Location", "Location of Incident"],
        "description": ["Description of Accident", "Accident Description", "Incident Description", "Description"],
    },
    "involvedParties": {
        "claimant": ["Claimant", "Insured Driver", "Driver Name", "Claimant Name"],
        "thirdParties": ["Third Parties", "Third Party", "Other Parties"],
        "contactDetails": ["Contact Details", "Contact Phone", "Contact Information", "Phone Number"],
    },
    "assetDetails": {
        "assetType": ["Asset Type", "Vehicle Type", "Body Type"],
        "assetId": ["Asset ID", "VIN", "Vehicle ID", "Vehicle Identification Number"],
        "estimatedDamage": ["Estimated Damage", "Damage Estimate", "Estimated Loss", "Estimated Repair Cost"],
    },
    "otherMandatoryFields": {
        "claimType": ["Claim Type", "Type of Claim"],
        "attachments": ["Attachments", "Supporting Documents", "Documents Attached"],
        "initialEstimate": ["Initial Estimate", "Initial Damage Estimate", "Initial Loss Estimate"],
    },
}

MANDATORY_FIELDS = [
    "policyNumber",
    "policyholderName",
    "date",
    "location",
    "description",
    "claimType",
    "estimatedDamage",
    "attachments",
    "initialEstimate",
]

_EMPTY_VALUES = {
    "", "n/a", "n.a.", "na", "none", "-", "--", "unknown", "tbd",
    "none provided", "not provided", "not available", "nil", "null",
}


def _all_label_variants():
    flat = {}
    for category, fields in FIELD_SCHEMA.items():
        for field_key, labels in fields.items():
            flat[field_key] = (category, labels)
    return flat


def _build_pattern(flat_fields):
    """Build one regex that captures each recognized label's value."""
    all_labels = []
    for field_key, (_category, labels) in flat_fields.items():
        for label in labels:
            all_labels.append((label, field_key))

    all_labels.sort(key=lambda pair: len(pair[0]), reverse=True)
    alternation = "|".join(re.escape(label) for label, _ in all_labels)

    pattern = re.compile(
        rf"(?P<label>{alternation})\s*(?:[:\-]|(?=\S))\s*"
        rf"(?P<value>.*?)"
        rf"(?=\n\s*(?:{alternation})\s*(?:[:\-]|\s)|\Z)",
        re.IGNORECASE | re.DOTALL,
    )
    label_to_field = {label.lower(): field_key for label, field_key in all_labels}
    return pattern, label_to_field


_FLAT_FIELDS = _all_label_variants()
_PATTERN, _LABEL_TO_FIELD = _build_pattern(_FLAT_FIELDS)


def _clean_value(value: str):
    """Normalize whitespace and turn known placeholders into None."""
    value = re.sub(r"\s*\n\s*", " ", value).strip()
    value = re.sub(r"[ \t]+", " ", value)
    if value.lower() in _EMPTY_VALUES:
        return None
    return value or None


def _read_docx_text(path: Path) -> str:
    """Extract paragraph and table text from a .docx Word document."""
    if Document is None:
        raise RuntimeError(
            "python-docx is required to read Word files. Install with: pip install python-docx"
        )

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                parts.append(row_text)
    return "\n".join(parts)


def _read_image_text(path: Path) -> str:
    """Extract text from a claim image using the bundled RapidOCR runtime."""
    if Image is None or RapidOCR is None:
        raise RuntimeError(
            "Image OCR dependencies are not installed. Install with: pip install Pillow rapidocr-onnxruntime"
        )

    image = Image.open(path).convert("RGB")
    ocr = RapidOCR()
    result, _elapsed = ocr(image)
    if not result:
        return ""

    lines = []
    for item in result:
        if len(item) >= 2 and item[1]:
            lines.append(str(item[1]))
    return "\n".join(lines)


def read_text_from_file(filepath: str) -> str:
    """Load raw text from a .txt, .pdf, .docx, or image FNOL document."""
    path = Path(filepath)
    suffix = path.suffix.lower()

    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")

    if suffix == ".pdf":
        if pdfplumber is None:
            raise RuntimeError(
                "pdfplumber is required to read PDF files. Install with: pip install pdfplumber"
            )
        text_parts = []
        with pdfplumber.open(str(path)) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                page_text = page.extract_text() or ""
                if page_text:
                    text_parts.append(f"[Page {page_number}]\n{page_text}")
        return "\n".join(text_parts)

    if suffix == ".docx":
        return _read_docx_text(path)

    if suffix in {".jpg", ".jpeg", ".png", ".webp"}:
        return _read_image_text(path)

    raise ValueError(f"Unsupported file type '{suffix}'. Supported FNOL inputs are TXT, PDF, DOCX, JPG, JPEG, PNG, and WEBP.")


def extract_fields(text: str) -> dict:
    """Extract known FNOL fields into the categorized FIELD_SCHEMA shape."""
    if not isinstance(text, str):
        raise TypeError("FNOL document text must be a string.")

    found = {}
    for match in _PATTERN.finditer(text):
        label = match.group("label").strip().lower()
        field_key = _LABEL_TO_FIELD.get(label)
        if not field_key:
            continue

        value = _clean_value(match.group("value"))
        if field_key not in found or found[field_key] is None:
            found[field_key] = value

    return {
        category: {field_key: found.get(field_key) for field_key in fields}
        for category, fields in FIELD_SCHEMA.items()
    }


def flatten_extracted(extracted: dict) -> dict:
    """Collapse categorized fields into a single {field_key: value} mapping."""
    flat = {}
    for fields in extracted.values():
        flat.update(fields)
    return flat
